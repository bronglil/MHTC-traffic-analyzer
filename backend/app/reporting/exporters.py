"""CSV / XLSX / JSON export of analysis results."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from collections.abc import Mapping, Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.vehicles import ALL_VEHICLE_TYPES

ZONE_COLUMNS = [
    "zone_id", "zone_name", "track_id", "vehicle_type", "direction",
    "first_time", "last_time", "first_frame", "last_frame", "frames_in_zone", "mean_confidence",
]
LINE_COLUMNS = ["line_id", "line_name", "track_id", "vehicle_type", "direction", "time", "frame"]


def _summary_tables(summary: Mapping[str, Any]) -> dict[str, tuple[list[str], list[list[Any]]]]:
    types = ALL_VEHICLE_TYPES
    area_rows = [[a["name"], a["total"], *[a["by_type"].get(t, 0) for t in types]] for a in summary["areas"]]
    line_rows = []
    for ln in summary["lines"]:
        for direction, n in sorted(ln["by_direction"].items()):
            line_rows.append([ln["name"], direction, n])
    time_rows = [
        [r["source"], r["name"], r["period"], r["total"], *[r["by_type"].get(t, 0) for t in types]]
        for r in summary["by_time"]
    ]
    return {
        "Summary by area": (["area", "total", *types], area_rows),
        "By area & type": (["area", "vehicle_type", "count"],
                           [[r["area"], r["vehicle_type"], r["count"]] for r in summary["by_area_type"]]),
        "By area & direction": (["area", "direction", "count"],
                                [[r["area"], r["direction"], r["count"]] for r in summary["by_area_direction"]]),
        "Line crossings by direction": (["line", "direction", "count"], line_rows),
        "By time period": (["source", "name", "period", "total", *types], time_rows),
    }


def _csv(columns: Sequence[str], rows: Sequence[Sequence[Any]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    w.writerows(rows)
    return buf.getvalue()


def _records(records: Sequence[Mapping[str, Any]], columns: Sequence[str]) -> list[list[Any]]:
    return [[r.get(c) for c in columns] for r in records]


def to_json(meta: Mapping[str, Any], summary: Mapping[str, Any], zone_counts, line_crossings) -> bytes:
    return json.dumps(
        {"analysis": meta, "summary": summary, "zone_counts": list(zone_counts), "line_crossings": list(line_crossings)},
        indent=2,
        default=str,
    ).encode()


def to_csv_zip(meta: Mapping[str, Any], summary: Mapping[str, Any], zone_counts, line_crossings) -> bytes:
    """A zip of CSVs, one per table (a single CSV cannot hold several tables)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for title, (cols, rows) in _summary_tables(summary).items():
            zf.writestr(title.lower().replace(" & ", "_").replace(" ", "_") + ".csv", _csv(cols, rows))
        zf.writestr("vehicle_counts.csv", _csv(ZONE_COLUMNS, _records(zone_counts, ZONE_COLUMNS)))
        zf.writestr("line_crossings.csv", _csv(LINE_COLUMNS, _records(line_crossings, LINE_COLUMNS)))
    return buf.getvalue()


def to_csv(zone_counts) -> bytes:
    """Single flat CSV of every unique vehicle-per-area count."""
    return _csv(ZONE_COLUMNS, _records(list(zone_counts), ZONE_COLUMNS)).encode()


def to_xlsx(meta: Mapping[str, Any], summary: Mapping[str, Any], zone_counts, line_crossings) -> bytes:
    wb = Workbook()
    info = wb.active
    info.title = "Analysis"
    for k, v in meta.items():
        info.append([k, json.dumps(v, default=str) if isinstance(v, dict | list) else v])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 60

    def sheet(title: str, cols: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
        ws = wb.create_sheet(title[:31])
        ws.append(list(cols))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(str(c)) + 2)

    for title, (cols, rows) in _summary_tables(summary).items():
        sheet(title, cols, rows)
    sheet("Vehicle counts", ZONE_COLUMNS, _records(list(zone_counts), ZONE_COLUMNS))
    sheet("Line crossing events", LINE_COLUMNS, _records(list(line_crossings), LINE_COLUMNS))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
