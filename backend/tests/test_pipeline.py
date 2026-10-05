import io
from pathlib import Path
import json
import zipfile

import pytest
from openpyxl import load_workbook

from app.analysis.analyzer import WHOLE_FRAME_ZONE_ID
from app.pipeline.engine import PipelineConfig, RegionDef, run_pipeline
from app.pipeline.tracker import IoUTracker, create_tracker
from app.reporting import exporters
from app.reporting.aggregate import summarize
from tests.conftest import FPS, FakeDetector

LEFT_HALF = RegionDef("left", "Left half", "polygon", [[0, 0], [0.5, 0], [0.5, 1], [0, 1]])
GATE = RegionDef("gate", "Centre gate", "line", [[0.5, 0], [0.5, 1]], "westbound", "eastbound")


def _run(video, tracker, **kw):
    cfg = PipelineConfig(vehicle_types=["car", "bus"], regions=[LEFT_HALF, GATE], **kw)
    return run_pipeline(str(video), cfg, FakeDetector(), tracker)


@pytest.mark.parametrize("tracker_kind", ["iou", "bytetrack", "botsort"])
def test_pipeline_counts(video_file, tracker_kind):
    progress, breakdowns = [], []
    cfg = PipelineConfig(vehicle_types=["car", "bus"], regions=[LEFT_HALF, GATE])
    result = run_pipeline(
        str(video_file), cfg, FakeDetector(), create_tracker(tracker_kind, FPS, 0.3),
        on_progress=lambda p, f, c: (progress.append(p), breakdowns.append(c)), progress_interval=0,
    )
    zones = {(c["zone_id"], c["vehicle_type"]): c for c in result.zone_counts}
    assert set(zones) == {(WHOLE_FRAME_ZONE_ID, "car"), (WHOLE_FRAME_ZONE_ID, "bus"), ("left", "car")}
    assert zones[(WHOLE_FRAME_ZONE_ID, "car")]["direction"] == "E"
    assert zones[(WHOLE_FRAME_ZONE_ID, "bus")]["direction"] == "S"
    assert [(c["line_id"], c["vehicle_type"], c["direction"]) for c in result.line_crossings] == [
        ("gate", "car", "eastbound")
    ]
    assert progress[-1] == 1.0
    assert breakdowns[-1]["gate"] == {"total": 1, "by_type": {"car": 1}, "by_direction": {"eastbound": 1}}
    assert breakdowns[-1]["left"]["by_type"] == {"car": 1}
    assert result.frames_processed == 60


def test_vehicle_type_selection_and_stride(video_file):
    cfg = PipelineConfig(vehicle_types=["bus"], frame_stride=2)
    result = run_pipeline(str(video_file), cfg, FakeDetector(), IoUTracker())
    assert [(c["zone_id"], c["vehicle_type"]) for c in result.zone_counts] == [(WHOLE_FRAME_ZONE_ID, "bus")]
    assert result.frames_processed == 30


def test_whole_frame_forced_without_polygons(video_file):
    cfg = PipelineConfig(vehicle_types=["car"], include_whole_frame=False, regions=[GATE])
    result = run_pipeline(str(video_file), cfg, FakeDetector(), IoUTracker())
    assert {c["zone_id"] for c in result.zone_counts} == {WHOLE_FRAME_ZONE_ID}


def test_annotated_video(video_file, tmp_path):
    out = tmp_path / "annotated.mp4"
    _run(video_file, IoUTracker(), annotated_video_path=str(out))
    import cv2

    cap = cv2.VideoCapture(str(out))
    assert int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) == 60
    cap.release()


def test_summary_and_exports(video_file):
    result = _run(video_file, IoUTracker())
    summary = summarize(result.zone_counts, result.line_crossings, time_bin_seconds=1)
    areas = {a["id"]: a for a in summary["areas"]}
    assert areas[WHOLE_FRAME_ZONE_ID]["total"] == 2
    assert areas["left"]["by_type"] == {"car": 1}
    assert summary["lines"][0]["by_direction"] == {"eastbound": 1}
    assert summary["lines"][0]["by_direction_type"] == {"eastbound": {"car": 1}}
    assert sum(r["total"] for r in summary["by_time"] if r["source"] == "area") == 3

    meta = {"analysis_id": "x"}
    data = json.loads(exporters.to_json(meta, summary, result.zone_counts, result.line_crossings))
    assert len(data["zone_counts"]) == 3

    csv_text = exporters.to_csv(result.zone_counts).decode()
    assert csv_text.splitlines()[0].startswith("zone_id,zone_name,track_id")
    assert len(csv_text.strip().splitlines()) == 4

    zf = zipfile.ZipFile(io.BytesIO(exporters.to_csv_zip(meta, summary, result.zone_counts, result.line_crossings)))
    assert "summary_by_area.csv" in zf.namelist()

    wb = load_workbook(io.BytesIO(exporters.to_xlsx(meta, summary, result.zone_counts, result.line_crossings)))
    assert "Summary by area" in wb.sheetnames
    assert wb["Vehicle counts"].max_row == 4


def test_stage_previews_and_pipeline_video(video_file, tmp_path):
    import cv2

    from app.pipeline.annotate import STAGES

    snapshots = []
    out = tmp_path / "pipeline.mp4"
    cfg = PipelineConfig(vehicle_types=["car", "bus"], regions=[LEFT_HALF, GATE], annotated_video_path=str(out),
                         annotated_video_layout="pipeline")
    run_pipeline(str(video_file), cfg, FakeDetector(), IoUTracker(),
                 on_stages=lambda st, i, t: snapshots.append((st, i)), stage_interval=0)
    assert len(snapshots) == 60
    stages, _ = snapshots[30]
    assert list(stages) == [k for k, _ in STAGES] + ["live"]
    assert all(img.shape == (360, 640, 3) for img in stages.values())
    # The detection panel differs from the raw frame (boxes drawn).
    assert (stages["detection"] != stages["frame"]).any()

    cap = cv2.VideoCapture(str(out))
    assert int(cap.get(cv2.CAP_PROP_FRAME_COUNT)) == 60
    assert (int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))) == (960, 412)
    cap.release()


def test_time_range_seeks_to_exact_frame():
    import numpy as np

    from app.pipeline.frames import iter_frames

    video = Path(__file__).parent / "data" / "videos" / "dual_carriageway.mp4"  # real H.264 with keyframe gaps
    sequential = {f.index: f.image for f in iter_frames(str(video)) if 140 <= f.index < 150}
    ranged = list(iter_frames(str(video), start_frame=140, end_frame=150))
    assert [f.index for f in ranged] == list(range(140, 150))
    for f in ranged:
        assert np.array_equal(f.image, sequential[f.index])
    assert [f.index for f in iter_frames(str(video), stride=3, start_frame=10, end_frame=20)] == [10, 13, 16, 19]


def test_time_range_limits_counting(video_file):
    # Car (left->right) and bus (top->bottom) are both visible for the whole 3 s clip;
    # analysing only 1.0-2.0 s processes 20 frames and stamps times on the video clock.
    progress = []
    cfg = PipelineConfig(vehicle_types=["car", "bus"], regions=[GATE], start_seconds=1.0, end_seconds=2.0)
    result = run_pipeline(str(video_file), cfg, FakeDetector(), IoUTracker(),
                          on_progress=lambda p, f, c: progress.append(p), progress_interval=0)
    assert result.frames_processed == 20
    assert all(1.0 <= c["first_time"] < 2.0 for c in result.zone_counts)
    assert progress[-1] == 1.0 and max(progress[:-1]) <= 0.99
    # The car crosses the centre gate at ~1.5 s, inside the window -> still counted.
    assert [lc["direction"] for lc in result.line_crossings] == ["eastbound"]
    late = run_pipeline(str(video_file), PipelineConfig(vehicle_types=["car"], regions=[GATE], start_seconds=2.0),
                        FakeDetector(), IoUTracker())
    assert late.line_crossings == []  # crossing happened before the window


def test_millisecond_timebase_reports_real_fps(video_file, monkeypatch):
    """Wikimedia Commons WebM clips make OpenCV report 1000 "fps" and a frame count
    in milliseconds; probe() must recover the real values from frame timestamps."""
    import cv2

    from app.pipeline import frames

    real_capture = cv2.VideoCapture

    class MillisecondTimebase:
        """Wraps a real capture but reports the bogus fps / count such files produce."""

        def __init__(self, *args):
            self._cap = real_capture(*args)

        def get(self, prop):
            if prop == cv2.CAP_PROP_FPS:
                return 1000.0
            if prop == cv2.CAP_PROP_FRAME_COUNT:
                return 3000.0  # 3.0 s expressed in ms, as OpenCV reports for such files
            return self._cap.get(prop)

        def __getattr__(self, name):
            return getattr(self._cap, name)

    monkeypatch.setattr(frames.cv2, "VideoCapture", MillisecondTimebase)
    info = frames.probe(str(video_file))
    assert info.fps == pytest.approx(FPS, rel=0.01)
    assert info.frame_count == 60
    assert info.duration == pytest.approx(3.0, rel=0.01)
